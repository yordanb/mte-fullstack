"""Proker Midlife Eng Booster (program kerja midlife engine pompa booster).

Sumber Excel: `midlife pompa booster.xlsx`, sheet `01. OVH DAN ML ENG`.
Berbeda dari Proker CAT (matriks unit x periode), sheet ini berupa satu baris
per unit dengan kolom milestone engine:

  NO | UNIT MODEL | PRODUCT | UNIT TYPE | C/N | HM NOW | INTERVAL MIDLIFE
  NEXT PLAN MIDLIFE / OVH (HM | DATE) | LAST (HM | DATE)
  OVH (HM | DATE) | MIDLIFE (HM | DATE) | LIFETIME (EG | MID) | REMARKS

Import memecahnya menjadi:
  * `booster_unit`  — identitas unit + HM NOW + interval + remarks
  * `booster_event` — riwayat event engine: OVH dan MIDLIFE (hm + tanggal)

Kolom turunan tidak disimpan, dihitung ulang saat baca:
  next_plan_hm = MAX(hm event) + interval
  hours_to_go  = next_plan_hm - HM NOW
  lifetime_eg  = HM NOW - HM OVH  (engine life sejak overhaul / ganti engine)
  lifetime_mid = |HM NOW - HM MIDLIFE|

Upsert event memakai (cn, kind, hm): upload ulang file yang sama hanya
memperbarui tanggal/catatan, sedangkan event dengan HM baru menambah riwayat.
"""
import datetime
import io
import json
import logging
import re

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.deps import get_current_user, require_menu
from app.db.session import get_db, SessionLocal

router = APIRouter(prefix="/v1/booster", tags=["booster"])
log = logging.getLogger("booster")

DEFAULT_INTERVAL = 20000            # INTERVAL MIDLIFE pada sheet
KINDS = ("OVH", "MIDLIFE")
BATCH = 500

UNIT_FIELDS = ("no", "unit_model", "product", "unit_type", "hm_now", "interval_hm", "remarks")


# --------------------------------------------------------------------------- util
def _txt(v) -> str | None:
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def _num(v) -> float | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _int(v) -> int | None:
    n = _num(v)
    return int(n) if n is not None else None


def _date(v) -> datetime.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        try:  # serial date Excel
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date()
        except (ValueError, OverflowError):
            return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%m/%d/%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# --------------------------------------------------------------------------- parser
def _find_sheet(wb):
    """Cari sheet yang punya header C/N (biasanya '01. OVH DAN ML ENG')."""
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(1, min(15, ws.max_row or 1) + 1):
            for c in range(1, min(40, ws.max_column or 1) + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip().upper() == "C/N":
                    return ws, r
    return None, None


def _parse(content: bytes, fname: str) -> tuple[dict, list[dict], list[dict]]:
    """Baca satu workbook -> (units{cn: row}, events[row], errors)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    try:
        ws, header_row = _find_sheet(wb)
        if ws is None:
            raise ValueError("kolom C/N tidak ditemukan — bukan format Proker Midlife Eng Booster")

        max_col = ws.max_column or 0
        max_row = ws.max_row or 0

        def cell(r: int, c: int):
            try:
                return ws.cell(row=r, column=c).value
            except (ValueError, IndexError):
                return None

        # label grup (baris header, termasuk sel merge) + sub-label (baris berikutnya)
        groups: dict[int, str] = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= header_row <= rng.max_row:
                val = _txt(cell(header_row, rng.min_col)) or ""
                for c in range(rng.min_col, rng.max_col + 1):
                    groups[c] = val
        subs: dict[int, str] = {}
        for c in range(1, max_col + 1):
            groups.setdefault(c, _txt(cell(header_row, c)) or "")
            subs[c] = (_txt(cell(header_row + 1, c)) or "").upper()

        def gl(c: int) -> str:
            return groups.get(c, "").strip().upper()

        def sl(c: int) -> str:
            return subs.get(c, "")

        cols: dict[str, int] = {}
        for c in range(1, max_col + 1):
            g, s = gl(c), sl(c)
            if g == "C/N":
                cols.setdefault("cn", c)
            elif g == "NO":
                cols.setdefault("no", c)
            elif g == "UNIT MODEL":
                cols.setdefault("unit_model", c)
            elif g == "PRODUCT":
                cols.setdefault("product", c)
            elif g == "UNIT TYPE":
                cols.setdefault("unit_type", c)
            elif g == "HM NOW":
                cols.setdefault("hm_now", c)
            elif g.startswith("INTERVAL"):
                cols.setdefault("interval_hm", c)
            elif g == "REMARKS":
                cols.setdefault("remarks", c)
            elif g.startswith("NEXT PLAN"):
                if s == "HM":
                    cols.setdefault("next_hm", c)
                elif s == "DATE":
                    cols.setdefault("next_date", c)
            elif g == "OVH":
                cols.setdefault("ovh_hm", c)
            elif g == "MIDLIFE":
                cols.setdefault("ml_hm", c)

        # kolom DATE tepat setelah OVH / MIDLIFE
        for key in ("ovh", "ml"):
            hm_col = cols.get(f"{key}_hm")
            if hm_col:
                nxt = hm_col + 1
                if gl(nxt) == "DATE" or sl(nxt) == "DATE":
                    cols[f"{key}_date"] = nxt

        if "cn" not in cols:
            raise ValueError("kolom C/N tidak ditemukan — bukan format Proker Midlife Eng Booster")
        if not any(k in cols for k in ("ovh_hm", "ml_hm", "hm_now")):
            raise ValueError("kolom OVH/MIDLIFE/HM NOW tidak ditemukan — format tidak dikenali")

        def val(r: int, key: str):
            c = cols.get(key)
            return cell(r, c) if c else None

        units: dict[str, dict] = {}
        events: list[dict] = []
        errors: list[dict] = []

        for r in range(header_row + 2, max_row + 1):
            cn = _txt(val(r, "cn"))
            if not cn:
                continue
            cn = cn.upper().replace(" ", "")
            interval = _num(val(r, "interval_hm"))
            if interval is None or interval <= 0:
                interval = float(DEFAULT_INTERVAL)
            units[cn] = {
                "cn": cn,
                "no": _int(val(r, "no")),
                "unit_model": _txt(val(r, "unit_model")),
                "product": _txt(val(r, "product")),
                "unit_type": _txt(val(r, "unit_type")),
                "hm_now": _num(val(r, "hm_now")),
                "interval_hm": interval,
                "remarks": _txt(val(r, "remarks")),
            }

            for kind, key in (("OVH", "ovh"), ("MIDLIFE", "ml")):
                raw_hm = val(r, f"{key}_hm")
                hm = _num(raw_hm)
                raw_dt = val(r, f"{key}_date")
                dt = _date(raw_dt)
                if hm is None:
                    if raw_dt is not None and dt is not None:
                        errors.append({"file": fname, "row": r, "cn": cn, "kind": kind,
                                       "error": f"HM {kind} kosong padahal ada tanggal {dt}"})
                    continue
                events.append({
                    "cn": cn, "kind": kind, "hm": hm, "service_date": dt, "note": None,
                })
                if raw_dt is not None and dt is None:
                    errors.append({"file": fname, "row": r, "cn": cn, "kind": kind,
                                   "error": f"tanggal {kind} tidak terbaca: {raw_dt!r}"})
        return units, events, errors
    finally:
        wb.close()


# --------------------------------------------------------------------------- import
_UNIT_UPSERT = text("""
INSERT INTO booster_unit
  (cn, no, unit_model, product, unit_type, hm_now, interval_hm, remarks, import_id)
VALUES
  (:cn, :no, :unit_model, :product, :unit_type, :hm_now, :interval_hm, :remarks, :import_id)
ON CONFLICT (cn) DO UPDATE SET
  no          = COALESCE(EXCLUDED.no, booster_unit.no),
  unit_model  = COALESCE(EXCLUDED.unit_model, booster_unit.unit_model),
  product     = COALESCE(EXCLUDED.product, booster_unit.product),
  unit_type   = COALESCE(EXCLUDED.unit_type, booster_unit.unit_type),
  hm_now      = COALESCE(EXCLUDED.hm_now, booster_unit.hm_now),
  interval_hm = EXCLUDED.interval_hm,
  remarks     = COALESCE(EXCLUDED.remarks, booster_unit.remarks),
  import_id   = EXCLUDED.import_id,
  updated_at  = now()
""")

_EVENT_UPSERT = text("""
INSERT INTO booster_event (cn, kind, hm, service_date, note, import_id)
VALUES (:cn, :kind, :hm, :service_date, :note, :import_id)
ON CONFLICT (cn, kind, hm) DO UPDATE SET
  service_date = COALESCE(EXCLUDED.service_date, booster_event.service_date),
  note         = COALESCE(EXCLUDED.note, booster_event.note),
  import_id    = EXCLUDED.import_id,
  updated_at   = now()
""")


async def _existing_events(db: AsyncSession, cns: list[str]) -> set[tuple[str, str, float]]:
    """Event (cn, kind, hm) yang sudah ada — untuk membedakan inserted vs updated."""
    if not cns:
        return set()
    rows = (await db.execute(text(
        "SELECT cn, kind, hm FROM booster_event WHERE cn = ANY(:cns)"), {"cns": cns}
    )).mappings().all()
    return {(r["cn"], r["kind"], float(r["hm"])) for r in rows}


async def _run(import_id: str, files: list[tuple[str, bytes]]):
    async with SessionLocal() as db:
        try:
            merged_units: dict[str, dict] = {}
            merged_events: dict[tuple[str, str, float], dict] = {}
            errors: list[dict] = []
            duplicates = 0

            for fname, content in files:
                try:
                    units, events, errs = _parse(content, fname)
                except Exception as e:  # noqa: BLE001
                    errors.append({"file": fname, "error": str(e)})
                    continue
                errors.extend(errs[:50])
                merged_units.update(units)
                for row in events:
                    key = (row["cn"], row["kind"], float(row["hm"]))
                    if key in merged_events:
                        duplicates += 1
                    merged_events[key] = row
                await db.execute(text(
                    "UPDATE imports SET total_rows=:t, processed_rows=:p WHERE id=:i"),
                    {"t": len(merged_events), "p": 0, "i": import_id})
                await db.commit()

            if not merged_units:
                raise ValueError("tidak ada baris unit yang terbaca (C/N kosong semua)")

            units = list(merged_units.values())
            for i in range(0, len(units), BATCH):
                await db.execute(_UNIT_UPSERT, [
                    {**u, "import_id": import_id} for u in units[i:i + BATCH]])

            events = list(merged_events.values())
            cns = sorted(merged_units)
            has = await _existing_events(db, cns)
            inserted = sum(1 for e in events
                           if (e["cn"], e["kind"], float(e["hm"])) not in has)
            updated = len(events) - inserted
            for i in range(0, len(events), BATCH):
                await db.execute(_EVENT_UPSERT, [
                    {**e, "import_id": import_id} for e in events[i:i + BATCH]])
                await db.execute(text(
                    "UPDATE imports SET processed_rows=:p, ok_rows=:o WHERE id=:i"),
                    {"p": min(i + BATCH, len(events)), "o": len(events), "i": import_id})
                await db.commit()

            ovh = sum(1 for e in events if e["kind"] == "OVH")
            summary = {
                "units": len(units), "events": len(events),
                "ovh": ovh, "midlife": len(events) - ovh,
                "inserted": inserted, "updated": updated,
                "duplicates": duplicates, "files": [f for f, _ in files],
                "total_rows": len(events), "errors": errors[:50],
            }
            await db.execute(text(
                "UPDATE imports SET status='COMMITTED', total_rows=:t, ok_rows=:o, "
                "fail_rows=:f, processed_rows=:p, errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"t": len(events), "o": len(events), "f": len(errors), "p": len(events),
                 "e": json.dumps(summary, default=str), "i": import_id})
            await db.commit()
            log.info("booster import %s: %s unit, %s event (%s baru), %s dup",
                     import_id, len(units), len(events), inserted, duplicates)
        except Exception as e:  # noqa: BLE001
            await db.rollback()
            await db.execute(text(
                "UPDATE imports SET status='FAILED', errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"e": json.dumps({"error": str(e)}, default=str), "i": import_id})
            await db.commit()
            log.exception("booster import gagal: %s", e)


@router.post("/import", dependencies=[Depends(require_menu("prokerbooster", "add"))], status_code=202)
async def import_xlsx(
    bg: BackgroundTasks,
    files: list[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """Upload 1..n file booster (.xlsx). Upsert unit (cn) + event (cn, kind, hm)."""
    if not files:
        raise HTTPException(400, "Tidak ada file")
    payload: list[tuple[str, bytes]] = []
    limit = settings.max_upload_mb * 1024 * 1024
    for f in files:
        content = await f.read()
        if len(content) > limit:
            raise HTTPException(413, f"{f.filename}: file terlalu besar")
        if not content:
            raise HTTPException(400, f"{f.filename}: file kosong")
        payload.append((f.filename or "upload.xlsx", content))
    label = ", ".join(name for name, _ in payload)
    res = await db.execute(text(
        "INSERT INTO imports(filename,sheet,status,total_rows,uploaded_by) "
        "VALUES (:fn,'BOOSTER','PROCESSING',0,:u) RETURNING id"),
        {"fn": label[:300], "u": user["username"]})
    import_id = str(res.scalar_one())
    await db.commit()
    bg.add_task(_run, import_id, payload)
    return {"import_id": import_id, "status": "PROCESSING", "files": len(payload)}


@router.get("/import/{import_id}", dependencies=[Depends(require_menu("prokerbooster", "view"))])
async def import_status(import_id: str, db: AsyncSession = Depends(get_db)):
    row = (await db.execute(text(
        "SELECT id, filename, status, total_rows, ok_rows, fail_rows, processed_rows, "
        "uploaded_by, created_at, errors FROM imports WHERE id=:i AND sheet='BOOSTER'"),
        {"i": import_id})).mappings().first()
    if not row:
        raise HTTPException(404, "import_id tidak ditemukan")
    return dict(row)


# --------------------------------------------------------------------------- agregat
_EVENTS_AGG = """
  SELECT cn,
         COUNT(*) AS event_count,
         MAX(hm) AS last_hm,
         MAX(hm) FILTER (WHERE kind = 'OVH') AS ovh_hm,
         MAX(service_date) FILTER (WHERE kind = 'OVH') AS ovh_date,
         MAX(hm) FILTER (WHERE kind = 'MIDLIFE') AS ml_hm,
         MAX(service_date) FILTER (WHERE kind = 'MIDLIFE') AS ml_date,
         (ARRAY_AGG(kind ORDER BY hm DESC, service_date DESC NULLS LAST))[1] AS last_kind,
         (ARRAY_AGG(service_date ORDER BY hm DESC, service_date DESC NULLS LAST))[1] AS last_date
  FROM booster_event
  GROUP BY cn
"""

_UNIT_SELECT = f"""
  WITH e AS ({_EVENTS_AGG})
  SELECT u.cn, u.no, u.unit_model, u.product, u.unit_type, u.hm_now, u.interval_hm,
         u.remarks,
         COALESCE(e.event_count, 0) AS event_count,
         e.ovh_hm, e.ovh_date, e.ml_hm, e.ml_date, e.last_hm, e.last_kind, e.last_date,
         COALESCE(e.last_hm, 0) + u.interval_hm AS next_plan_hm,
         CASE WHEN u.hm_now IS NULL THEN NULL
              ELSE (COALESCE(e.last_hm, 0) + u.interval_hm) - u.hm_now END AS hours_to_go,
         CASE WHEN u.hm_now IS NULL THEN NULL
              WHEN e.ovh_hm IS NULL THEN u.hm_now
              WHEN u.hm_now < e.ovh_hm THEN u.hm_now
              ELSE u.hm_now - e.ovh_hm END AS lifetime_eg,
         CASE WHEN u.hm_now IS NULL OR e.ml_hm IS NULL THEN NULL
              ELSE ABS(u.hm_now - e.ml_hm) END AS lifetime_mid
  FROM booster_unit u LEFT JOIN e ON e.cn = u.cn
"""

_SORT = {
    "cn": "u.cn", "no": "u.no", "unit_model": "u.unit_model", "product": "u.product",
    "hm_now": "u.hm_now", "next_plan_hm": "next_plan_hm", "hours_to_go": "hours_to_go",
    "last_date": "last_date", "event_count": "event_count",
}


@router.get("/filters", dependencies=[Depends(require_menu("prokerbooster", "view"))])
async def get_filters(db: AsyncSession = Depends(get_db)):
    async def distinct(col: str) -> list[str]:
        return list((await db.execute(text(
            f"SELECT DISTINCT {col} FROM booster_unit "
            f"WHERE {col} IS NOT NULL AND {col} <> '' ORDER BY 1"))).scalars().all())

    span = (await db.execute(text(
        "SELECT MIN(service_date) AS min_date, MAX(service_date) AS max_date "
        "FROM booster_event"))).mappings().first()
    return {
        "product": await distinct("product"),
        "model": await distinct("unit_model"),
        "type": await distinct("unit_type"),
        "interval": DEFAULT_INTERVAL,
        "min_date": str(span["min_date"]) if span and span["min_date"] else None,
        "max_date": str(span["max_date"]) if span and span["max_date"] else None,
    }


@router.get("/stats", dependencies=[Depends(require_menu("prokerbooster", "view"))])
async def get_stats(db: AsyncSession = Depends(get_db)):
    units = (await db.execute(text("SELECT COUNT(*) FROM booster_unit"))).scalar() or 0
    ev = (await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE kind = 'OVH') AS ovh,
               COUNT(*) FILTER (WHERE kind = 'MIDLIFE') AS midlife,
               COUNT(DISTINCT cn) AS unit_event,
               MAX(service_date) AS last_date
        FROM booster_event
    """))).mappings().first()
    due = (await db.execute(text(f"""
        WITH e AS ({_EVENTS_AGG})
        SELECT COUNT(*) FROM booster_unit u LEFT JOIN e ON e.cn = u.cn
        WHERE u.hm_now IS NOT NULL
          AND (COALESCE(e.last_hm, 0) + u.interval_hm - u.hm_now) <= 0
    """))).scalar() or 0
    return {
        "units": units, "total": ev["total"], "ovh": ev["ovh"], "midlife": ev["midlife"],
        "unit_event": ev["unit_event"], "due": due,
        "last_date": str(ev["last_date"]) if ev["last_date"] else None,
        "interval": DEFAULT_INTERVAL,
    }


@router.get("/units", dependencies=[Depends(require_menu("prokerbooster", "view"))])
async def list_units(
    search: str | None = None,
    product: str | None = None,
    model: str | None = None,
    due: bool | None = None,
    sort: str = Query("cn"),
    order: str = Query("asc", pattern="^(asc|desc)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    conds, p = [], {"lim": page_size, "off": (page - 1) * page_size}
    if search:
        conds.append("(u.cn ILIKE :s OR u.unit_model ILIKE :s OR u.product ILIKE :s "
                     "OR u.unit_type ILIKE :s OR u.remarks ILIKE :s)")
        p["s"] = f"%{search.strip()}%"
    if product:
        conds.append("u.product = :prod")
        p["prod"] = product
    if model:
        conds.append("u.unit_model = :model")
        p["model"] = model
    if due is not None:
        conds.append("(:due = FALSE OR (u.hm_now IS NOT NULL AND "
                     "(COALESCE(e.last_hm, 0) + u.interval_hm - u.hm_now) <= 0))")
        p["due"] = due
    where = f"WHERE {' AND '.join(conds)}" if conds else ""
    col = _SORT.get(sort, "u.cn")
    total = (await db.execute(text(f"""
        WITH e AS ({_EVENTS_AGG})
        SELECT COUNT(*) FROM booster_unit u LEFT JOIN e ON e.cn = u.cn {where}
    """), p)).scalar() or 0
    rows = (await db.execute(text(f"""
        {_UNIT_SELECT}
        {where}
        ORDER BY {col} {order.upper()} NULLS LAST, u.cn
        LIMIT :lim OFFSET :off
    """), p)).mappings().all()
    return {"total": total, "page": page, "page_size": page_size,
            "data": [dict(r) for r in rows]}


@router.get("/units/{cn}", dependencies=[Depends(require_menu("prokerbooster", "view"))])
async def unit_detail(cn: str, db: AsyncSession = Depends(get_db)):
    key = (cn or "").strip().upper().replace(" ", "")
    unit = (await db.execute(text(f"{_UNIT_SELECT} WHERE u.cn = :c"),
                             {"c": key})).mappings().first()
    if not unit:
        raise HTTPException(404, f"{key} tidak ditemukan di Proker Midlife Eng Booster")
    rows = (await db.execute(text("""
        SELECT * FROM booster_event WHERE cn = :c
        ORDER BY hm DESC, service_date DESC NULLS LAST
    """), {"c": key})).mappings().all()
    events = [dict(r) for r in rows]
    return {
        "unit": dict(unit),
        "events": events,
        "stats": {
            "events": len(events),
            "ovh": sum(1 for e in events if e["kind"] == "OVH"),
            "midlife": sum(1 for e in events if e["kind"] == "MIDLIFE"),
            "hm_now": unit["hm_now"],
            "interval_hm": unit["interval_hm"],
            "last_hm": unit["last_hm"],
            "last_kind": unit["last_kind"],
            "last_date": unit["last_date"],
            "next_plan_hm": unit["next_plan_hm"],
            "hours_to_go": unit["hours_to_go"],
            "lifetime_eg": unit["lifetime_eg"],
            "lifetime_mid": unit["lifetime_mid"],
            "due": (unit["hours_to_go"] is not None and float(unit["hours_to_go"]) <= 0),
        },
    }


# --------------------------------------------------------------------------- tulis
@router.post("/units", dependencies=[Depends(require_menu("prokerbooster", "add"))], status_code=201)
async def create_unit(body: dict, db: AsyncSession = Depends(get_db)):
    cn = (_txt(body.get("cn")) or "").upper().replace(" ", "")
    if not cn:
        raise HTTPException(400, "cn wajib diisi")
    if (await db.execute(text("SELECT 1 FROM booster_unit WHERE cn=:c"), {"c": cn})).scalar():
        raise HTTPException(409, f"{cn} sudah ada")
    interval = _num(body.get("interval_hm"))
    if interval is None or interval <= 0:
        interval = float(DEFAULT_INTERVAL)
    row = (await db.execute(text("""
        INSERT INTO booster_unit
          (cn, no, unit_model, product, unit_type, hm_now, interval_hm, remarks)
        VALUES (:cn, :no, :unit_model, :product, :unit_type, :hm_now, :interval_hm, :remarks)
        RETURNING *
    """), {
        "cn": cn, "no": _int(body.get("no")),
        "unit_model": _txt(body.get("unit_model")), "product": _txt(body.get("product")),
        "unit_type": _txt(body.get("unit_type")), "hm_now": _num(body.get("hm_now")),
        "interval_hm": interval, "remarks": _txt(body.get("remarks")),
    })).mappings().first()
    await db.commit()
    return dict(row)


@router.patch("/units/{cn}", dependencies=[Depends(require_menu("prokerbooster", "edit"))])
async def update_unit(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    key = cn.strip().upper()
    sets, params = [], {"c": key}
    for f in UNIT_FIELDS:
        if f not in body:
            continue
        if f == "hm_now":
            params[f] = _num(body[f])
        elif f == "interval_hm":
            val = _num(body[f])
            params[f] = val if val and val > 0 else float(DEFAULT_INTERVAL)
        elif f == "no":
            params[f] = _int(body[f])
        else:
            params[f] = _txt(body[f])
        sets.append(f"{f}=:{f}")
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    res = await db.execute(text(
        f"UPDATE booster_unit SET {', '.join(sets)} WHERE cn=:c"), params)
    if res.rowcount == 0:
        raise HTTPException(404, f"{key} tidak ditemukan")
    await db.commit()
    row = (await db.execute(text("SELECT * FROM booster_unit WHERE cn=:c"),
                            {"c": key})).mappings().first()
    return dict(row)


@router.post("/units/{cn}/events",
             dependencies=[Depends(require_menu("prokerbooster", "add"))], status_code=201)
async def create_event(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    key = cn.strip().upper()
    if not (await db.execute(text("SELECT 1 FROM booster_unit WHERE cn=:c"), {"c": key})).scalar():
        raise HTTPException(404, f"{key} belum ada di master unit")
    kind = (_txt(body.get("kind")) or "").upper()
    if kind not in KINDS:
        raise HTTPException(400, "jenis event harus OVH atau MIDLIFE")
    hm = _num(body.get("hm"))
    if hm is None:
        raise HTTPException(400, "HM wajib diisi angka")
    dup = (await db.execute(text(
        "SELECT 1 FROM booster_event WHERE cn=:c AND kind=:k AND hm=:h"),
        {"c": key, "k": kind, "h": hm})).scalar()
    if dup:
        raise HTTPException(409, f"{kind} {hm:g} jam untuk {key} sudah ada")
    row = (await db.execute(text("""
        INSERT INTO booster_event (cn, kind, hm, service_date, note)
        VALUES (:cn, :kind, :hm, :d, :note) RETURNING *
    """), {
        "cn": key, "kind": kind, "hm": hm,
        "d": _date(body.get("service_date")), "note": _txt(body.get("note")),
    })).mappings().first()
    await db.commit()
    return dict(row)


@router.patch("/events/{event_id}", dependencies=[Depends(require_menu("prokerbooster", "edit"))])
async def update_event(event_id: int, body: dict, db: AsyncSession = Depends(get_db)):
    cur = (await db.execute(text("SELECT * FROM booster_event WHERE id=:i"),
                            {"i": event_id})).mappings().first()
    if not cur:
        raise HTTPException(404, "baris event tidak ditemukan")

    sets, params = [], {"i": event_id}
    kind = cur["kind"]
    if "kind" in body:
        kind = (_txt(body["kind"]) or "").upper()
        if kind not in KINDS:
            raise HTTPException(400, "jenis event harus OVH atau MIDLIFE")
        sets.append("kind=:kind")
        params["kind"] = kind
    hm = cur["hm"] if "hm" not in body else _num(body["hm"])
    if hm is None:
        raise HTTPException(400, "HM wajib diisi angka")
    if "hm" in body:
        sets.append("hm=:hm")
        params["hm"] = hm
    if "kind" in body or "hm" in body:
        dup = (await db.execute(text(
            "SELECT 1 FROM booster_event WHERE cn=:c AND kind=:k AND hm=:h AND id<>:i"),
            {"c": cur["cn"], "k": kind, "h": hm, "i": event_id})).scalar()
        if dup:
            raise HTTPException(409, f"{kind} {float(hm):g} jam untuk {cur['cn']} sudah ada")
    if "service_date" in body:
        sets.append("service_date=:d")
        params["d"] = _date(body["service_date"])
    if "note" in body:
        sets.append("note=:note")
        params["note"] = _txt(body["note"])
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    row = (await db.execute(text(
        f"UPDATE booster_event SET {', '.join(sets)} WHERE id=:i RETURNING *"),
        params)).mappings().first()
    await db.commit()
    return dict(row)


@router.delete("/events/{event_id}",
               dependencies=[Depends(require_menu("prokerbooster", "delete"))])
async def delete_event(event_id: int, db: AsyncSession = Depends(get_db)):
    res = await db.execute(text("DELETE FROM booster_event WHERE id=:i"), {"i": event_id})
    if res.rowcount == 0:
        raise HTTPException(404, "baris event tidak ditemukan")
    await db.commit()
    return {"deleted": res.rowcount}
