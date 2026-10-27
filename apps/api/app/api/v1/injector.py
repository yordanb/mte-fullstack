"""Proker Midlife Injector Primer (replace injector primer 16.000 jam).

Sumber Excel: `midlife injector primer.xlsx`, sheet `04. INJECTOR DAN TURBO`.
Satu baris per unit dengan tiga blok penggantian:

  NO | UNIT MODEL | PRODUCT | UNIT TYPE | C/N | HM NOW | LIFETIME (OVH | MID)
  PENGGANTIAN 1 : 12 kolom posisi injector 1..12
  PENGGANTIAN 2 : 12 kolom posisi injector 1..12
  PENGGANTIAN 3 : 12 kolom posisi injector 1..12

Tiap sel = HM saat injector posisi tsb diganti, teks "PLAN" bila baru
direncanakan, atau kosong bila belum ada rencana.

Import memecahnya menjadi:
  * `injector_unit`        — identitas unit + HM NOW + interval
  * `injector_replacement` — satu baris per (penggantian ke-, nomor injector)

Kolom turunan tidak disimpan, dihitung ulang saat baca:
  umur per posisi = HM NOW - HM penggantian terakhir posisi tsb
  umur unit       = HM NOW - HM penggantian terakhir unit
  next_plan_hm    = HM penggantian terakhir + interval
  hours_to_go     = next_plan_hm - HM NOW

Kolom LIFETIME (G/H/I) pada sheet adalah formula turunan / VLOOKUP ke
workbook lain, jadi tidak diimpor.
"""
import io
import json
import logging
import re

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.deps import get_current_user, require_menu
from app.db.session import get_db, SessionLocal

router = APIRouter(prefix="/v1/injector", tags=["injector"])
log = logging.getLogger("injector")

DEFAULT_INTERVAL = 16000           # REPLACE INJECTOR PRIMER (16.000 HOURS)
POSITIONS = 12                     # jumlah injector per unit
STATUSES = ("DONE", "PLAN")
BATCH = 500

UNIT_FIELDS = ("no", "unit_model", "product", "unit_type", "hm_now", "interval_hm")

_GROUP_RE = re.compile(r"PENGGANTIAN\s*(\d+)", re.IGNORECASE)


# --------------------------------------------------------------------------- util
def _txt(v) -> str | None:
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def _num(v) -> float | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _int(v) -> int | None:
    n = _num(v)
    return int(n) if n is not None else None


def _cn(v) -> str | None:
    s = _txt(v)
    return s.upper().replace(" ", "") if s else None


# --------------------------------------------------------------------------- parser
def _find_sheet(wb):
    """Cari sheet yang punya header C/N (biasanya '04. INJECTOR DAN TURBO')."""
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(1, min(15, ws.max_row or 1) + 1):
            for c in range(1, min(60, ws.max_column or 1) + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip().upper() == "C/N":
                    return ws, r
    return None, None


def _parse(content: bytes, fname: str) -> tuple[dict, list[dict], list[dict]]:
    """Baca satu workbook -> (units{cn: row}, replacements[row], errors)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    try:
        ws, header_row = _find_sheet(wb)
        if ws is None:
            raise ValueError("kolom C/N tidak ditemukan — bukan format Proker Midlife Injector Primer")

        max_col = ws.max_column or 0
        max_row = ws.max_row or 0

        def cell(r: int, c: int):
            try:
                return ws.cell(row=r, column=c).value
            except (ValueError, IndexError):
                return None

        # label grup pada baris header (termasuk sel merge) + sub-label di bawahnya
        groups: dict[int, str] = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= header_row <= rng.max_row:
                val = _txt(cell(header_row, rng.min_col)) or ""
                for c in range(rng.min_col, rng.max_col + 1):
                    groups[c] = val
        for c in range(1, max_col + 1):
            groups.setdefault(c, _txt(cell(header_row, c)) or "")

        cols: dict[str, int] = {}
        repl_cols: list[tuple[int, int, int]] = []   # (seq, position, column)
        for c in range(1, max_col + 1):
            g = groups.get(c, "").strip().upper()
            sub = _txt(cell(header_row + 1, c))
            if g == "C/N":
                cols.setdefault("cn", c)
            elif g == "NO":
                cols.setdefault("no", c)
            elif g == "UNIT MODEL":
                cols.setdefault("unit_model", c)
            elif g == "PRODUCT":
                cols.setdefault("product", c)
            elif g == "UNIT TYPE":
                cols.setdefault("unit_type", c)
            elif g == "HM NOW":
                cols.setdefault("hm_now", c)
            else:
                m = _GROUP_RE.search(g)
                if m:
                    pos = _int(sub)
                    if pos is None:
                        pos = _int(cell(header_row + 2, c))
                    if pos is not None:
                        repl_cols.append((int(m.group(1)), pos, c))

        if "cn" not in cols:
            raise ValueError("kolom C/N tidak ditemukan — bukan format Proker Midlife Injector Primer")
        if "hm_now" not in cols:
            raise ValueError("kolom HM NOW tidak ditemukan — format tidak dikenali")
        if not repl_cols:
            raise ValueError("kolom PENGGANTIAN tidak ditemukan — format tidak dikenali")

        def val(r: int, key: str):
            c = cols.get(key)
            return cell(r, c) if c else None

        units: dict[str, dict] = {}
        repls: list[dict] = []
        errors: list[dict] = []

        for r in range(header_row + 2, max_row + 1):
            cn = _cn(val(r, "cn"))
            if not cn:
                continue
            units[cn] = {
                "cn": cn,
                "no": _int(val(r, "no")),
                "unit_model": _txt(val(r, "unit_model")),
                "product": _txt(val(r, "product")),
                "unit_type": _txt(val(r, "unit_type")),
                "hm_now": _num(val(r, "hm_now")),
                "interval_hm": float(DEFAULT_INTERVAL),
            }

            for seq, pos, c in repl_cols:
                raw = cell(r, c)
                if raw is None:
                    continue
                if isinstance(raw, str):
                    s = raw.replace("\xa0", " ").strip()
                    if not s:
                        continue
                    if s.upper() == "PLAN":
                        repls.append({
                            "cn": cn, "seq": seq, "position": pos,
                            "hm": None, "status": "PLAN", "note": None,
                        })
                        continue
                    num = _num(s)
                    if num is None:
                        errors.append({
                            "file": fname, "row": r, "cn": cn, "seq": seq,
                            "position": pos, "error": f"nilai tidak dikenali: {s!r}",
                        })
                        continue
                    repls.append({
                        "cn": cn, "seq": seq, "position": pos,
                        "hm": num, "status": "DONE", "note": None,
                    })
                else:
                    num = _num(raw)
                    if num is None:
                        errors.append({
                            "file": fname, "row": r, "cn": cn, "seq": seq,
                            "position": pos, "error": f"angka tidak terbaca: {raw!r}",
                        })
                        continue
                    repls.append({
                        "cn": cn, "seq": seq, "position": pos,
                        "hm": num, "status": "DONE", "note": None,
                    })
        return units, repls, errors
    finally:
        wb.close()


# --------------------------------------------------------------------------- import
_UNIT_UPSERT = text("""
INSERT INTO injector_unit
  (cn, no, unit_model, product, unit_type, hm_now, interval_hm, import_id)
VALUES
  (:cn, :no, :unit_model, :product, :unit_type, :hm_now, :interval_hm, :import_id)
ON CONFLICT (cn) DO UPDATE SET
  no          = COALESCE(EXCLUDED.no, injector_unit.no),
  unit_model  = COALESCE(EXCLUDED.unit_model, injector_unit.unit_model),
  product     = COALESCE(EXCLUDED.product, injector_unit.product),
  unit_type   = COALESCE(EXCLUDED.unit_type, injector_unit.unit_type),
  hm_now      = COALESCE(EXCLUDED.hm_now, injector_unit.hm_now),
  interval_hm = EXCLUDED.interval_hm,
  import_id   = EXCLUDED.import_id,
  updated_at  = now()
""")

_REPL_UPSERT = text("""
INSERT INTO injector_replacement (cn, seq, position, hm, status, note, import_id)
VALUES (:cn, :seq, :position, :hm, :status, :note, :import_id)
ON CONFLICT (cn, seq, position) DO UPDATE SET
  hm        = EXCLUDED.hm,
  status    = EXCLUDED.status,
  note      = COALESCE(EXCLUDED.note, injector_replacement.note),
  import_id = EXCLUDED.import_id,
  updated_at = now()
""")


async def _existing(db: AsyncSession, cns: list[str]) -> set[tuple[str, int, int]]:
    if not cns:
        return set()
    rows = (await db.execute(text(
        "SELECT cn, seq, position FROM injector_replacement WHERE cn = ANY(:cns)"),
        {"cns": cns})).mappings().all()
    return {(r["cn"], int(r["seq"]), int(r["position"])) for r in rows}


async def _run(import_id: str, files: list[tuple[str, bytes]]):
    async with SessionLocal() as db:
        try:
            merged_units: dict[str, dict] = {}
            merged: dict[tuple[str, int, int], dict] = {}
            errors: list[dict] = []
            duplicates = 0

            for fname, content in files:
                try:
                    units, repls, errs = _parse(content, fname)
                except Exception as e:  # noqa: BLE001
                    errors.append({"file": fname, "error": str(e)})
                    continue
                errors.extend(errs[:50])
                merged_units.update(units)
                for row in repls:
                    key = (row["cn"], int(row["seq"]), int(row["position"]))
                    if key in merged:
                        duplicates += 1
                    merged[key] = row
                await db.execute(text(
                    "UPDATE imports SET total_rows=:t, processed_rows=:p WHERE id=:i"),
                    {"t": len(merged), "p": 0, "i": import_id})
                await db.commit()

            if not merged_units:
                raise ValueError("tidak ada baris unit yang terbaca (C/N kosong semua)")

            units = list(merged_units.values())
            for i in range(0, len(units), BATCH):
                await db.execute(_UNIT_UPSERT, [
                    {**u, "import_id": import_id} for u in units[i:i + BATCH]])

            rows = list(merged.values())
            cns = sorted(merged_units)
            has = await _existing(db, cns)
            inserted = sum(1 for r in rows
                           if (r["cn"], int(r["seq"]), int(r["position"])) not in has)
            updated = len(rows) - inserted
            for i in range(0, len(rows), BATCH):
                await db.execute(_REPL_UPSERT, [
                    {**r, "import_id": import_id} for r in rows[i:i + BATCH]])
                await db.execute(text(
                    "UPDATE imports SET processed_rows=:p, ok_rows=:o WHERE id=:i"),
                    {"p": min(i + BATCH, len(rows)), "o": len(rows), "i": import_id})
                await db.commit()

            done = sum(1 for r in rows if r["status"] == "DONE")
            summary = {
                "units": len(units), "replacements": len(rows),
                "done": done, "plan": len(rows) - done,
                "unit_with_repl": len({r["cn"] for r in rows}),
                "inserted": inserted, "updated": updated,
                "duplicates": duplicates, "files": [f for f, _ in files],
                "total_rows": len(rows), "errors": errors[:50],
            }
            await db.execute(text(
                "UPDATE imports SET status='COMMITTED', total_rows=:t, ok_rows=:o, "
                "fail_rows=:f, processed_rows=:p, errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"t": len(rows), "o": len(rows), "f": len(errors), "p": len(rows),
                 "e": json.dumps(summary, default=str), "i": import_id})
            await db.commit()
            log.info("injector import %s: %s unit, %s penggantian (%s baru), %s dup",
                     import_id, len(units), len(rows), inserted, duplicates)
        except Exception as e:  # noqa: BLE001
            await db.rollback()
            await db.execute(text(
                "UPDATE imports SET status='FAILED', errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"e": json.dumps({"error": str(e)}, default=str), "i": import_id})
            await db.commit()
            log.exception("injector import gagal: %s", e)


@router.post("/import", dependencies=[Depends(require_menu("prokerinjector", "add"))], status_code=202)
async def import_xlsx(
    bg: BackgroundTasks,
    files: list[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """Upload 1..n file injector (.xlsx). Upsert unit (cn) + baris (cn, seq, position)."""
    if not files:
        raise HTTPException(400, "Tidak ada file")
    payload: list[tuple[str, bytes]] = []
    limit = settings.max_upload_mb * 1024 * 1024
    for f in files:
        content = await f.read()
        if len(content) > limit:
            raise HTTPException(413, f"{f.filename}: file terlalu besar")
        if not content:
            raise HTTPException(400, f"{f.filename}: file kosong")
        payload.append((f.filename or "upload.xlsx", content))
    label = ", ".join(name for name, _ in payload)
    res = await db.execute(text(
        "INSERT INTO imports(filename,sheet,status,total_rows,uploaded_by) "
        "VALUES (:fn,'INJECTOR','PROCESSING',0,:u) RETURNING id"),
        {"fn": label[:300], "u": user["username"]})
    import_id = str(res.scalar_one())
    await db.commit()
    bg.add_task(_run, import_id, payload)
    return {"import_id": import_id, "status": "PROCESSING", "files": len(payload)}


@router.get("/import/{import_id}", dependencies=[Depends(require_menu("prokerinjector", "view"))])
async def import_status(import_id: str, db: AsyncSession = Depends(get_db)):
    row = (await db.execute(text(
        "SELECT id, filename, status, total_rows, ok_rows, fail_rows, processed_rows, "
        "uploaded_by, created_at, errors FROM imports WHERE id=:i AND sheet='INJECTOR'"),
        {"i": import_id})).mappings().first()
    if not row:
        raise HTTPException(404, "import_id tidak ditemukan")
    return dict(row)


# --------------------------------------------------------------------------- agregat
_REPL_AGG = """
  SELECT cn,
         COUNT(*) AS repl_count,
         COUNT(*) FILTER (WHERE status = 'DONE') AS done_count,
         COUNT(*) FILTER (WHERE status = 'PLAN') AS plan_count,
         MAX(hm) FILTER (WHERE status = 'DONE') AS last_hm,
         MAX(seq) AS last_seq
  FROM injector_replacement
  GROUP BY cn
"""

_UNIT_SELECT = f"""
  WITH r AS ({_REPL_AGG})
  SELECT u.cn, u.no, u.unit_model, u.product, u.unit_type, u.hm_now, u.interval_hm,
         COALESCE(r.repl_count, 0) AS repl_count,
         COALESCE(r.done_count, 0) AS done_count,
         COALESCE(r.plan_count, 0) AS plan_count,
         r.last_hm, r.last_seq,
         CASE WHEN u.hm_now IS NULL OR r.last_hm IS NULL THEN NULL
              ELSE u.hm_now - r.last_hm END AS age_hm,
         COALESCE(r.last_hm, 0) + u.interval_hm AS next_plan_hm,
         CASE WHEN u.hm_now IS NULL OR r.last_hm IS NULL THEN NULL
              ELSE (r.last_hm + u.interval_hm) - u.hm_now END AS hours_to_go
  FROM injector_unit u LEFT JOIN r ON r.cn = u.cn
"""

_SORT = {
    "cn": "u.cn", "no": "u.no", "unit_model": "u.unit_model", "product": "u.product",
    "hm_now": "u.hm_now", "age_hm": "age_hm", "hours_to_go": "hours_to_go",
    "next_plan_hm": "next_plan_hm", "repl_count": "repl_count",
}


@router.get("/filters", dependencies=[Depends(require_menu("prokerinjector", "view"))])
async def get_filters(db: AsyncSession = Depends(get_db)):
    async def distinct(col: str) -> list[str]:
        return list((await db.execute(text(
            f"SELECT DISTINCT {col} FROM injector_unit "
            f"WHERE {col} IS NOT NULL AND {col} <> '' ORDER BY 1"))).scalars().all())

    return {
        "product": await distinct("product"),
        "model": await distinct("unit_model"),
        "type": await distinct("unit_type"),
        "interval": DEFAULT_INTERVAL,
        "positions": POSITIONS,
    }


@router.get("/stats", dependencies=[Depends(require_menu("prokerinjector", "view"))])
async def get_stats(db: AsyncSession = Depends(get_db)):
    units = (await db.execute(text("SELECT COUNT(*) FROM injector_unit"))).scalar() or 0
    agg = (await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE status = 'DONE') AS done,
               COUNT(*) FILTER (WHERE status = 'PLAN') AS plan,
               COUNT(DISTINCT cn) AS unit_with_repl
        FROM injector_replacement
    """))).mappings().first()
    due = (await db.execute(text(f"""
        WITH r AS ({_REPL_AGG})
        SELECT COUNT(*) FROM injector_unit u LEFT JOIN r ON r.cn = u.cn
        WHERE u.hm_now IS NOT NULL AND r.last_hm IS NOT NULL
          AND (u.hm_now - r.last_hm) >= u.interval_hm
    """))).scalar() or 0
    return {
        "units": units, "total": agg["total"], "done": agg["done"], "plan": agg["plan"],
        "unit_with_repl": agg["unit_with_repl"], "due": due,
        "interval": DEFAULT_INTERVAL, "positions": POSITIONS,
    }


@router.get("/units", dependencies=[Depends(require_menu("prokerinjector", "view"))])
async def list_units(
    search: str | None = None,
    product: str | None = None,
    model: str | None = None,
    due: bool | None = None,
    sort: str = Query("cn"),
    order: str = Query("asc", pattern="^(asc|desc)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    conds, p = [], {"lim": page_size, "off": (page - 1) * page_size}
    if search:
        conds.append("(u.cn ILIKE :s OR u.unit_model ILIKE :s OR u.product ILIKE :s "
                     "OR u.unit_type ILIKE :s)")
        p["s"] = f"%{search.strip()}%"
    if product:
        conds.append("u.product = :prod")
        p["prod"] = product
    if model:
        conds.append("u.unit_model = :model")
        p["model"] = model
    if due is not None:
        conds.append("(:due = FALSE OR (u.hm_now IS NOT NULL AND r.last_hm IS NOT NULL "
                     "AND (u.hm_now - r.last_hm) >= u.interval_hm))")
        p["due"] = due
    where = f"WHERE {' AND '.join(conds)}" if conds else ""
    col = _SORT.get(sort, "u.cn")
    total = (await db.execute(text(f"""
        WITH r AS ({_REPL_AGG})
        SELECT COUNT(*) FROM injector_unit u LEFT JOIN r ON r.cn = u.cn {where}
    """), p)).scalar() or 0
    rows = (await db.execute(text(f"""
        {_UNIT_SELECT}
        {where}
        ORDER BY {col} {order.upper()} NULLS LAST, u.cn
        LIMIT :lim OFFSET :off
    """), p)).mappings().all()
    return {"total": total, "page": page, "page_size": page_size,
            "data": [dict(r) for r in rows]}


@router.get("/units/{cn}", dependencies=[Depends(require_menu("prokerinjector", "view"))])
async def unit_detail(cn: str, db: AsyncSession = Depends(get_db)):
    key = _cn(cn) or ""
    unit = (await db.execute(text(f"{_UNIT_SELECT} WHERE u.cn = :c"),
                             {"c": key})).mappings().first()
    if not unit:
        raise HTTPException(404, f"{key} tidak ditemukan di Proker Midlife Injector Primer")

    rows = (await db.execute(text("""
        SELECT * FROM injector_replacement WHERE cn = :c
        ORDER BY seq DESC, position
    """), {"c": key})).mappings().all()
    repls = [dict(r) for r in rows]

    # grid posisi: pakai baris dengan seq terbesar untuk tiap posisi
    hm_now = unit["hm_now"]
    by_pos: dict[int, dict] = {}
    for r in sorted(repls, key=lambda x: (int(x["seq"]), int(x["position"]))):
        by_pos[int(r["position"])] = r
    positions = []
    for pos in range(1, POSITIONS + 1):
        row = by_pos.get(pos)
        hm = float(row["hm"]) if row and row["hm"] is not None else None
        positions.append({
            "position": pos,
            "seq": row["seq"] if row else None,
            "hm": hm,
            "status": row["status"] if row else None,
            "age": None if hm is None or hm_now is None else float(hm_now) - hm,
        })

    done = [r for r in repls if r["status"] == "DONE"]
    return {
        "unit": dict(unit),
        "replacements": repls,
        "positions": positions,
        "stats": {
            "replacements": len(repls),
            "done": len(done),
            "plan": len(repls) - len(done),
            "hm_now": hm_now,
            "interval_hm": unit["interval_hm"],
            "last_hm": unit["last_hm"],
            "last_seq": unit["last_seq"],
            "age_hm": unit["age_hm"],
            "next_plan_hm": unit["next_plan_hm"],
            "hours_to_go": unit["hours_to_go"],
            "due": (unit["age_hm"] is not None
                    and float(unit["age_hm"]) >= float(unit["interval_hm"])),
        },
    }


# --------------------------------------------------------------------------- tulis
@router.post("/units", dependencies=[Depends(require_menu("prokerinjector", "add"))], status_code=201)
async def create_unit(body: dict, db: AsyncSession = Depends(get_db)):
    cn = _cn(body.get("cn"))
    if not cn:
        raise HTTPException(400, "cn wajib diisi")
    if (await db.execute(text("SELECT 1 FROM injector_unit WHERE cn=:c"), {"c": cn})).scalar():
        raise HTTPException(409, f"{cn} sudah ada")
    interval = _num(body.get("interval_hm"))
    if interval is None or interval <= 0:
        interval = float(DEFAULT_INTERVAL)
    row = (await db.execute(text("""
        INSERT INTO injector_unit
          (cn, no, unit_model, product, unit_type, hm_now, interval_hm)
        VALUES (:cn, :no, :unit_model, :product, :unit_type, :hm_now, :interval_hm)
        RETURNING *
    """), {
        "cn": cn, "no": _int(body.get("no")),
        "unit_model": _txt(body.get("unit_model")), "product": _txt(body.get("product")),
        "unit_type": _txt(body.get("unit_type")), "hm_now": _num(body.get("hm_now")),
        "interval_hm": interval,
    })).mappings().first()
    await db.commit()
    return dict(row)


@router.patch("/units/{cn}", dependencies=[Depends(require_menu("prokerinjector", "edit"))])
async def update_unit(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    key = _cn(cn) or ""
    sets, params = [], {"c": key}
    for f in UNIT_FIELDS:
        if f not in body:
            continue
        if f == "hm_now":
            params[f] = _num(body[f])
        elif f == "interval_hm":
            v = _num(body[f])
            params[f] = v if v and v > 0 else float(DEFAULT_INTERVAL)
        elif f == "no":
            params[f] = _int(body[f])
        else:
            params[f] = _txt(body[f])
        sets.append(f"{f}=:{f}")
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    res = await db.execute(text(
        f"UPDATE injector_unit SET {', '.join(sets)} WHERE cn=:c"), params)
    if res.rowcount == 0:
        raise HTTPException(404, f"{key} tidak ditemukan")
    await db.commit()
    row = (await db.execute(text("SELECT * FROM injector_unit WHERE cn=:c"),
                            {"c": key})).mappings().first()
    return dict(row)


@router.post("/units/{cn}/replacements",
             dependencies=[Depends(require_menu("prokerinjector", "add"))], status_code=201)
async def create_replacements(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    """Catat satu penggantian untuk beberapa posisi sekaligus."""
    key = _cn(cn) or ""
    if not (await db.execute(text("SELECT 1 FROM injector_unit WHERE cn=:c"), {"c": key})).scalar():
        raise HTTPException(404, f"{key} belum ada di master unit")

    seq = _int(body.get("seq"))
    if seq is None or not (1 <= seq <= 9):
        raise HTTPException(400, "seq (penggantian ke-) harus 1..9")

    raw_pos = body.get("positions")
    if not isinstance(raw_pos, list) or not raw_pos:
        raise HTTPException(400, "positions wajib diisi (array nomor injector)")
    positions = sorted({int(p) for p in raw_pos if str(p).strip().isdigit()})
    if not positions or any(p < 1 or p > 24 for p in positions):
        raise HTTPException(400, "positions harus berisi nomor injector 1..24")

    status = (_txt(body.get("status")) or "DONE").upper()
    if status not in STATUSES:
        raise HTTPException(400, "status harus DONE atau PLAN")
    hm = _num(body.get("hm"))
    if status == "DONE" and hm is None:
        raise HTTPException(400, "HM wajib diisi untuk status DONE")
    if status == "PLAN":
        hm = None

    note = _txt(body.get("note"))
    await db.execute(text("""
        INSERT INTO injector_replacement (cn, seq, position, hm, status, note)
        VALUES (:cn, :seq, :position, :hm, :status, :note)
        ON CONFLICT (cn, seq, position) DO UPDATE SET
          hm = EXCLUDED.hm, status = EXCLUDED.status,
          note = COALESCE(EXCLUDED.note, injector_replacement.note),
          updated_at = now()
    """), [{"cn": key, "seq": seq, "position": p, "hm": hm,
            "status": status, "note": note} for p in positions])
    rows = (await db.execute(text("""
        SELECT * FROM injector_replacement
        WHERE cn = :c AND seq = :s AND position = ANY(:ps)
        ORDER BY position
    """), {"c": key, "s": seq, "ps": positions})).mappings().all()
    await db.commit()
    return {"cn": key, "seq": seq, "created": len(rows), "data": [dict(r) for r in rows]}


@router.patch("/replacements/{repl_id}",
              dependencies=[Depends(require_menu("prokerinjector", "edit"))])
async def update_replacement(repl_id: int, body: dict, db: AsyncSession = Depends(get_db)):
    cur = (await db.execute(text("SELECT * FROM injector_replacement WHERE id=:i"),
                            {"i": repl_id})).mappings().first()
    if not cur:
        raise HTTPException(404, "baris penggantian tidak ditemukan")

    sets, params = [], {"i": repl_id}
    status = cur["status"]
    if "status" in body:
        status = (_txt(body["status"]) or "").upper()
        if status not in STATUSES:
            raise HTTPException(400, "status harus DONE atau PLAN")
        sets.append("status=:status")
        params["status"] = status
    if "hm" in body:
        hm = _num(body["hm"])
        sets.append("hm=:hm")
        params["hm"] = hm
    elif status == "PLAN":
        sets.append("hm=:hm")
        params["hm"] = None
    elif status == "DONE" and cur["hm"] is None:
        raise HTTPException(400, "HM wajib diisi untuk status DONE")
    if "seq" in body:
        seq = _int(body["seq"])
        if seq is None or not (1 <= seq <= 9):
            raise HTTPException(400, "seq harus 1..9")
        dup = (await db.execute(text(
            "SELECT 1 FROM injector_replacement WHERE cn=:c AND seq=:s AND position=:p AND id<>:i"),
            {"c": cur["cn"], "s": seq, "p": cur["position"], "i": repl_id})).scalar()
        if dup:
            raise HTTPException(409, f"penggantian {seq} posisi {cur['position']} sudah ada")
        sets.append("seq=:seq")
        params["seq"] = seq
    if "position" in body:
        pos = _int(body["position"])
        if pos is None or not (1 <= pos <= 24):
            raise HTTPException(400, "position harus 1..24")
        dup = (await db.execute(text(
            "SELECT 1 FROM injector_replacement WHERE cn=:c AND seq=:s AND position=:p AND id<>:i"),
            {"c": cur["cn"], "s": cur["seq"], "p": pos, "i": repl_id})).scalar()
        if dup:
            raise HTTPException(409, f"penggantian {cur['seq']} posisi {pos} sudah ada")
        sets.append("position=:position")
        params["position"] = pos
    if "note" in body:
        sets.append("note=:note")
        params["note"] = _txt(body["note"])
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    row = (await db.execute(text(
        f"UPDATE injector_replacement SET {', '.join(sets)} WHERE id=:i RETURNING *"),
        params)).mappings().first()
    await db.commit()
    return dict(row)


@router.delete("/replacements/{repl_id}",
               dependencies=[Depends(require_menu("prokerinjector", "delete"))])
async def delete_replacement(repl_id: int, db: AsyncSession = Depends(get_db)):
    res = await db.execute(text("DELETE FROM injector_replacement WHERE id=:i"), {"i": repl_id})
    if res.rowcount == 0:
        raise HTTPException(404, "baris penggantian tidak ditemukan")
    await db.commit()
    return {"deleted": res.rowcount}


@router.delete("/units/{cn}/replacements/{seq}",
               dependencies=[Depends(require_menu("prokerinjector", "delete"))])
async def delete_replacement_group(cn: str, seq: int, db: AsyncSession = Depends(get_db)):
    """Hapus satu penggantian penuh (semua posisi pada seq tsb)."""
    res = await db.execute(text(
        "DELETE FROM injector_replacement WHERE cn=:c AND seq=:s"),
        {"c": _cn(cn) or "", "s": seq})
    if res.rowcount == 0:
        raise HTTPException(404, "penggantian tidak ditemukan")
    await db.commit()
    return {"deleted": res.rowcount}
