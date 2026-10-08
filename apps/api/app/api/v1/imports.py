import io
import openpyxl
from fastapi import APIRouter, Depends, UploadFile, File, Query, HTTPException, BackgroundTasks
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.core.deps import get_current_user, require_role
from app.db.session import get_db, SessionLocal
from app.services.excel_map import HEADER_MAP, HEADER_ROW, DATA_START_ROW, SHEET, coerce

router = APIRouter(prefix="/v1/imports", tags=["imports"])

ALL_COLS = list(HEADER_MAP.values())  # 116 kolom, urutan sesuai HEADER_MAP
BATCH = 2000  # upsert per batch agar 36rb baris tidak 36rb round-trip

def _header_idx(ws):
    """Index header yang toleran spasi ekstra. Dipakai dry-run dan commit."""
    header = list(next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)))
    idx = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        key = str(h).strip()
        if key and key not in idx:
            idx[key] = i
    unknown = [h for h in idx if h not in HEADER_MAP]
    if unknown:
        raise HTTPException(400, f"Header tak dikenal: {unknown}")
    missing = [h for h in HEADER_MAP if h not in idx]
    if missing:
        raise HTTPException(400, f"Header hilang: {missing}")
    return idx


def _open_sheet(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise HTTPException(400, f"Sheet {SHEET} tidak ditemukan, ada: {wb.sheetnames}")
    return wb[SHEET]


def _read_full(content: bytes):
    ws = _open_sheet(content)
    idx = _header_idx(ws)
    rows, errors = [], []
    for rno, r in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW):
        if all(v is None for v in r):
            continue
        try:
            clean = {col: coerce(col, r[idx[h]]) for h, col in HEADER_MAP.items()}
            if clean["lab_no"] is None:
                raise ValueError("Lab No kosong")
            rows.append(clean)
        except Exception as e:
            errors.append({"row": rno, "error": str(e)})
    return rows, errors

UPSERT_SQL = None

def _upsert_sql() -> str:
    global UPSERT_SQL
    if UPSERT_SQL:
        return UPSERT_SQL
    cols = ALL_COLS + ["import_id"]
    col_list = ", ".join([f'"{c}"' if c == "condition" else c for c in cols])
    val_list = ", ".join(f":{c}" for c in cols)
    sets = ", ".join(
        f'"{c}"=EXCLUDED."{c}"' if c == "condition" else f"{c}=EXCLUDED.{c}"
        for c in cols if c not in ("lab_no",))
    UPSERT_SQL = (f"INSERT INTO oil_lab_result({col_list}) VALUES ({val_list}) "
                  f"ON CONFLICT (lab_no) DO UPDATE SET {sets}")
    return UPSERT_SQL

async def _run_commit(import_id: str, content: bytes):
    """Background job penuh: parse (progres per 5000 baris) + upsert batch 2000."""
    async with SessionLocal() as db:
        try:
            ws = _open_sheet(content)
            idx = _header_idx(ws)
            rows, errors, n = [], [], 0
            for rno, r in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW):
                if all(v is None for v in r):
                    continue
                n += 1
                try:
                    clean = {col: coerce(col, r[idx[h]]) for h, col in HEADER_MAP.items()}
                    if clean["lab_no"] is None:
                        raise ValueError("Lab No kosong")
                    rows.append(clean)
                except Exception as e:
                    errors.append({"row": rno, "error": str(e)})
                if n % 5000 == 0:
                    await db.execute(text(
                        "UPDATE imports SET total_rows=:t, processed_rows=:p WHERE id=:i"),
                        {"t": n, "p": n, "i": import_id})
                    await db.commit()
            import json as _json
            await db.execute(text(
                "UPDATE imports SET total_rows=:t, ok_rows=:o, fail_rows=:f, "
                "processed_rows=0, errors=:e WHERE id=:i"),
                {"t": len(rows) + len(errors), "o": len(rows),
                 "f": len(errors), "e": _json.dumps(errors[:50]), "i": import_id})
            await db.commit()
            sql = _upsert_sql()
            done = 0
            for n in range(0, len(rows), BATCH):
                batch = [{**r, "import_id": import_id} for r in rows[n:n + BATCH]]
                await db.execute(text(sql), batch)
                done += len(batch)
                await db.execute(text(
                    "UPDATE imports SET processed_rows=:p WHERE id=:i"),
                    {"p": done, "i": import_id})
                await db.commit()
            await db.execute(text(
                "UPDATE imports SET status='COMMITTED', processed_rows=:p WHERE id=:i"),
                {"p": done, "i": import_id})
            await db.commit()
            try:
                await db.execute(text("REFRESH MATERIALIZED VIEW CONCURRENTLY mv_latest_status"))
            except Exception:
                await db.rollback()
                await db.execute(text("REFRESH MATERIALIZED VIEW mv_latest_status"))
            await db.commit()
        except Exception:
            await db.rollback()
            await db.execute(text(
                "UPDATE imports SET status='FAILED' WHERE id=:i"), {"i": import_id})
            await db.commit()
            raise

@router.post("", dependencies=[Depends(require_role("inputer", "admin"))], status_code=202)
async def upload_excel(
    bg: BackgroundTasks,
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "File terlalu besar")
    if dry_run:
        rows, errors = _read_full(content)
        return {"total": len(rows) + len(errors), "ok": len(rows), "fail": len(errors),
                "errors": errors[:50],
                "preview": [{k: rows[0][k] for k in ("lab_no", "vesselid", "unit_id", "sample_date", "condition")} ] if rows else []}
    # commit: validasi header dulu agar gagal-cepat (400), bukan ok=0/fail=N.
    # Langsung 202, parse + upsert semua di background agar progres terpantau
    _header_idx(_open_sheet(content))
    res = await db.execute(text(
        "INSERT INTO imports(filename,status,total_rows,ok_rows,fail_rows,processed_rows,uploaded_by) "
        "VALUES (:fn,'PROCESSING',0,0,0,0,:u) RETURNING id"),
        {"fn": file.filename, "u": user["username"]})
    import_id = str(res.scalar_one())
    await db.commit()
    bg.add_task(_run_commit, import_id, content)
    return {"import_id": import_id, "status": "PROCESSING"}

@router.get("/latest")
async def latest_import(db: AsyncSession = Depends(get_db),
                        user=Depends(get_current_user)):
    """Upload report oli terakhir (DBR dikecualikan via sheet)."""
    row = (await db.execute(text(
        "SELECT id, filename, status, total_rows, ok_rows, fail_rows, "
        "processed_rows, uploaded_by, created_at FROM imports "
        "WHERE sheet != 'DBR' "
        "ORDER BY created_at DESC LIMIT 1"))).mappings().first()
    return dict(row) if row else {}

@router.get("/{import_id}")
async def import_status(import_id: str, db: AsyncSession = Depends(get_db),
                        user=Depends(get_current_user)):
    """Progres import untuk polling progress bar."""
    row = (await db.execute(text(
        "SELECT id, filename, status, total_rows, ok_rows, fail_rows, "
        "processed_rows, uploaded_by, created_at FROM imports WHERE id=:i"),
        {"i": import_id})).mappings().first()
    if not row:
        raise HTTPException(404, "import_id tidak ditemukan")
    return dict(row)
