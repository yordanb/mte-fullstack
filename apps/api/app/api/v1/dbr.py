import datetime
import io
import openpyxl
from fastapi import APIRouter, Depends, UploadFile, File, Query, HTTPException, BackgroundTasks
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.core.deps import get_current_user, require_role
from app.db.session import get_db, SessionLocal
from app.services.dbr_map import HEADER_MAP, HEADER_ROW, DATA_START_ROW, SHEET, coerce

router = APIRouter(prefix="/v1/dbr", tags=["dbr"])

ALL_COLS = list(HEADER_MAP.values())
BATCH = 2000

def _upsert_sql() -> str:
    cols = ALL_COLS + ["import_id"]
    sets = ", ".join(f"{c}=EXCLUDED.{c}" for c in cols if c != "id")
    return (f"INSERT INTO dbr_records({', '.join(cols)}) "
            f"VALUES ({', '.join(':' + c for c in cols)}) "
            f"ON CONFLICT (date, cn, start_breakdown) DO UPDATE SET {sets}")

async def _run(import_id: str, content: bytes):
    async with SessionLocal() as db:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            if SHEET not in wb.sheetnames:
                raise ValueError(f"Sheet {SHEET} tidak ditemukan")
            ws = wb[SHEET]
            header = list(next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)))
            idx = {h: i for i, h in enumerate(header) if h}
            rows, errors, n = [], [], 0
            for rno, r in enumerate(ws.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW):
                if all(v is None for v in r):
                    continue
                n += 1
                try:
                    rows.append({col: coerce(col, r[idx[h]]) for h, col in HEADER_MAP.items()})
                except Exception as e:
                    errors.append({"row": rno, "error": str(e)})
                if n % 5000 == 0:
                    await db.execute(text(
                        "UPDATE imports SET total_rows=:t, processed_rows=:p WHERE id=:i"),
                        {"t": n, "p": n, "i": import_id})
                    await db.commit()
            import json as _json
            await db.execute(text(
                "UPDATE imports SET total_rows=:t, ok_rows=:o, fail_rows=:f, "
                "processed_rows=0, errors=:e WHERE id=:i"),
                {"t": len(rows) + len(errors), "o": len(rows),
                 "f": len(errors), "e": _json.dumps(errors[:50]), "i": import_id})
            await db.commit()
            sql = _upsert_sql()
            done = 0
            for i in range(0, len(rows), BATCH):
                batch = [{**r, "import_id": import_id} for r in rows[i:i + BATCH]]
                await db.execute(text(sql), batch)
                done += len(batch)
                await db.execute(text(
                    "UPDATE imports SET processed_rows=:p WHERE id=:i"),
                    {"p": done, "i": import_id})
                await db.commit()
            await db.execute(text(
                "UPDATE imports SET status='COMMITTED', processed_rows=:p WHERE id=:i"),
                {"p": done, "i": import_id})
            await db.commit()
        except Exception:
            await db.rollback()
            await db.execute(text("UPDATE imports SET status='FAILED' WHERE id=:i"),
                             {"i": import_id})
            await db.commit()
            raise

@router.post("/imports", dependencies=[Depends(require_role("inputer", "admin"))], status_code=202)
async def upload_dbr(bg: BackgroundTasks, file: UploadFile = File(...),
                     db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "File terlalu besar")
    res = await db.execute(text(
        "INSERT INTO imports(filename,sheet,status,total_rows,uploaded_by) "
        "VALUES (:fn,'DBR','PROCESSING',0,:u) RETURNING id"),
        {"fn": file.filename, "u": user["username"]})
    import_id = str(res.scalar_one())
    await db.commit()
    bg.add_task(_run, import_id, content)
    return {"import_id": import_id, "status": "PROCESSING"}

@router.get("/codes")
async def codes(db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    rows = (await db.execute(text(
        "SELECT DISTINCT code FROM dbr_records WHERE code IS NOT NULL ORDER BY code"))
    ).scalars().all()
    return {"data": list(rows)}

@router.get("/records")
async def records(
    date_from: datetime.date | None = None,
    date_to: datetime.date | None = None,
    cn: str | None = None,
    prefix: str | None = Query(None, min_length=2, max_length=2),
    code: str | None = None,
    exclude_continue: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """Tabel DBR: default 30 hari terakhir kalau tanpa filter tanggal."""
    if date_from is None and date_to is None:
        date_to = datetime.date.today()
        date_from = date_to - datetime.timedelta(days=30)
    conds, p = [], {"lim": page_size, "off": (page - 1) * page_size}
    if date_from:
        conds.append("date >= :df"); p["df"] = date_from
    if date_to:
        conds.append("date <= :dt"); p["dt"] = date_to
    if cn:
        conds.append("cn = UPPER(:cn)"); p["cn"] = cn
    if prefix:
        conds.append("cn_prefix = UPPER(:pfx)"); p["pfx"] = prefix
    if code == "__EMPTY__":
        conds.append("code IS NULL")
    elif code:
        conds.append("code = UPPER(:code)"); p["code"] = code
    if exclude_continue:
        conds.append("(action IS NULL OR UPPER(action) != 'CONTINUE')")
    where = f"WHERE {' AND '.join(conds)}" if conds else ""
    total = (await db.execute(text(f"SELECT count(*) FROM dbr_records {where}"), p)).scalar()
    rows = (await db.execute(text(
        f"SELECT * FROM dbr_records {where} ORDER BY date DESC, cn LIMIT :lim OFFSET :off"),
        p)).mappings().all()
    return {"total": total, "page": page, "page_size": page_size, "data": [dict(r) for r in rows]}

@router.get("/stats")
async def stats(
    date_from: datetime.date | None = None,
    date_to: datetime.date | None = None,
    granularity: str = Query("week", pattern="^(day|week|month)$"),
    prefix: str | None = Query(None, min_length=2, max_length=2),
    code: str | None = None,
    exclude_continue: bool = True,
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """Agregat frekuensi breakdown per periode + Pareto trouble/section/code.

    exclude_continue=True (default): baris ACTION=CONTINUE (carry-over
    perbaikan multi-hari) tidak dihitung sebagai kejadian baru.
    """
    if date_from is None and date_to is None:
        date_to = datetime.date.today()
        date_from = date_to - datetime.timedelta(days=90)
    conds, p = [], {}
    if date_from:
        conds.append("date >= :df"); p["df"] = date_from
    if date_to:
        conds.append("date <= :dt"); p["dt"] = date_to
    if prefix:
        conds.append("cn_prefix = UPPER(:pfx)"); p["pfx"] = prefix
    if code == "__EMPTY__":
        conds.append("code IS NULL")
    elif code:
        conds.append("code = UPPER(:code)"); p["code"] = code
    if exclude_continue:
        conds.append("(action IS NULL OR UPPER(action) != 'CONTINUE')")
    where = f"WHERE {' AND '.join(conds)}" if conds else ""
    trunc = {"day": "day", "week": "week", "month": "month"}[granularity]
    series = (await db.execute(text(
        f"SELECT date_trunc('{trunc}', date)::date AS period, cn_prefix AS prefix, count(*) AS n "
        f"FROM dbr_records {where} GROUP BY 1, 2 ORDER BY 1, 2"), p)
    ).mappings().all()

    async def top(col: str, exclude_null: bool):
        w = f"{where} AND {col} IS NOT NULL AND TRIM({col}) <> ''" if where else \
            f"WHERE {col} IS NOT NULL AND TRIM({col}) <> ''"
        if not exclude_null:
            w = where
        q = (f"SELECT {col} AS k, count(*) AS v FROM dbr_records {w} "
             f"GROUP BY 1 ORDER BY 2 DESC LIMIT 10")
        return [dict(r) for r in (await db.execute(text(q), p)).mappings().all()]

    summ = (await db.execute(text(
        f"SELECT count(*) AS total, count(DISTINCT cn) AS units, "
        f"count(DISTINCT date) AS days, "
        f"count(*) FILTER (WHERE code IS NULL) AS empty_code "
        f"FROM dbr_records {where}"), p)).mappings().first()
    excluded = 0
    if exclude_continue:
        w2 = f"{where} AND UPPER(action) = 'CONTINUE'" if where else \
            "WHERE UPPER(action) = 'CONTINUE'"
        excluded = (await db.execute(
            text(f"SELECT count(*) FROM dbr_records {w2}"), p)).scalar() or 0
    return {"granularity": granularity, "exclude_continue": exclude_continue,
            "excluded_continue": excluded,
            "series": [dict(r) for r in series],
            "top_trouble": await top("trouble", True),
            "top_section": await top("section", True),
            "top_code": await top("code", True),
            "top_cn": await top("cn", True),
            "summary": dict(summ) if summ else {}}

