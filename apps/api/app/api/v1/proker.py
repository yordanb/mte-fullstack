"""Proker (program kerja pompa primer 2.000 hours).

Sumber Excel: `gmc pompa primer.xlsx`, sheet `02. GMC` — matriks 49 unit x N periode.
Tiap periode memakai 3 kolom berurutan: `HM | DATE | MEKANIK`.

Import memecah matriks menjadi:
  * `proker_unit`      — identitas unit (cn, model, product, type, hm_now)
  * `proker_execution` — riwayat eksekusi (period, hm, tanggal, mekanik, status)

Upsert berbasis (cn, period), jadi meng-upload file yang sama/lebih baru akan
memperbarui baris lama tanpa menghapus data di luar file.
"""
import datetime
import io
import json
import logging
import re

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.deps import get_current_user, require_menu
from app.db.session import get_db, SessionLocal

router = APIRouter(prefix="/v1/proker", tags=["proker"])
log = logging.getLogger("proker")

INTERVAL_HM = 2000                       # judul sheet: "GMC POMPA PRIMER (2.000 HOURS)"
STATUSES = ("DONE", "PLAN", "LATE", "MISSED")
_SEQ = {"pertama": 1, "kedua": 2, "ketiga": 3, "keempat": 4}
_YEAR_RE = re.compile(r"(?:19|20)\d{2}")
BATCH = 500

UNIT_FIELDS = ("no", "unit_model", "product", "unit_type", "hm_now")
EXEC_FIELDS = ("period", "hm", "service_date", "mekanik", "status")


# --------------------------------------------------------------------------- util
def _txt(v) -> str | None:
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def _num(v) -> float | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _int(v) -> int | None:
    n = _num(v)
    return int(n) if n is not None else None


def _date(v) -> datetime.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        try:  # serial date Excel
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date()
        except (ValueError, OverflowError):
            return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%m/%d/%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _period_parts(label: str) -> tuple[int | None, int | None]:
    """'2025 - Pertama' -> (2025, 1); '2022' -> (2022, None)."""
    m = _YEAR_RE.search(label or "")
    year = int(m.group(0)) if m else None
    low = (label or "").lower()
    seq = next((n for k, n in _SEQ.items() if k in low), None)
    return year, seq


def _norm_status(raw, hm: float | None) -> str:
    s = (_txt(raw) or "").upper()
    if s in STATUSES:
        return s
    return "DONE" if hm is not None else "PLAN"


# --------------------------------------------------------------------------- parser
def _find_sheet(wb):
    """Cari sheet yang punya header C/N (biasanya '02. GMC')."""
    for sn in wb.sheetnames:
        ws = wb[sn]
        max_r = min(15, ws.max_row or 1)
        max_c = min(40, ws.max_column or 1)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip().upper() == "C/N":
                    return ws, r, c
    return None, None, None


def _parse(content: bytes, fname: str) -> tuple[dict, list[dict], list[dict]]:
    """Baca satu workbook -> (units{cn: row}, executions[row], errors)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    try:
        ws, header_row, cn_col = _find_sheet(wb)
        if ws is None:
            raise ValueError("kolom C/N tidak ditemukan — bukan format Proker")

        def cell(r: int, c: int):
            try:
                return ws.cell(row=r, column=c).value
            except (ValueError, IndexError):
                return None

        def head(c: int) -> str:
            v = cell(header_row, c)
            return str(v).strip().upper() if v is not None else ""

        # kolom identitas (A..F)
        cols: dict[str, int] = {}
        for c in range(1, (ws.max_column or 0) + 1):
            h = head(c)
            if h in ("NO", "UNIT MODEL", "PRODUCT", "UNIT TYPE", "HM NOW"):
                cols.setdefault(h, c)

        # label periode dari sel merge di baris header
        labels: dict[int, str] = {}
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= header_row <= rng.max_row:
                labels[rng.min_col] = (_txt(cell(header_row, rng.min_col)) or "", rng.max_col)

        def label_at(c: int) -> str:
            for start, (val, end) in labels.items():
                if start <= c <= end:
                    return val
            return _txt(cell(header_row, c)) or ""

        sub_row = header_row + 1

        def sub(c: int) -> str:
            v = cell(sub_row, c)
            return str(v).strip().upper() if v is not None else ""

        # grup periode = HM | DATE | MEKANIK berurutan
        groups: list[tuple[str, int, int, int]] = []
        for c in range(cn_col + 1, (ws.max_column or 0) - 1):
            if sub(c) == "HM" and sub(c + 1) == "DATE" and sub(c + 2) == "MEKANIK":
                groups.append((label_at(c) or f"Periode {len(groups) + 1}", c, c + 1, c + 2))
        if not groups:
            raise ValueError("kolom periode (HM | DATE | MEKANIK) tidak ditemukan")

        def col(name: str):
            return cols.get(name)

        units: dict[str, dict] = {}
        execs: list[dict] = []
        errors: list[dict] = []

        for r in range(header_row + 2, (ws.max_row or 0) + 1):
            cn = _txt(cell(r, cn_col))
            if not cn:
                continue
            cn = cn.upper().replace(" ", "")
            units[cn] = {
                "cn": cn,
                "no": _int(cell(r, col("NO"))) if col("NO") else None,
                "unit_model": _txt(cell(r, col("UNIT MODEL"))) if col("UNIT MODEL") else None,
                "product": _txt(cell(r, col("PRODUCT"))) if col("PRODUCT") else None,
                "unit_type": _txt(cell(r, col("UNIT TYPE"))) if col("UNIT TYPE") else None,
                "hm_now": _num(cell(r, col("HM NOW"))) if col("HM NOW") else None,
            }

            for label, ch, cd, cm in groups:
                hm_raw = cell(r, ch)
                date_raw = cell(r, cd)
                mek = _txt(cell(r, cm))
                hm = _num(hm_raw)
                dt = _date(date_raw)
                raw_status = hm_raw.strip().upper() if isinstance(hm_raw, str) else ""
                # "PLAN"/"LATE" bukan angka tapi tetap baris proker yang harus disimpan.
                if hm is None and dt is None and not mek \
                        and raw_status not in ("PLAN", "LATE", "RENCANA", "PLANNING"):
                    continue
                if mek and mek.lower().startswith("missed"):
                    status = "MISSED"
                elif raw_status in ("PLAN", "RENCANA", "PLANNING"):
                    status = "PLAN"
                elif raw_status == "LATE":
                    status = "LATE"
                else:
                    status = "DONE"
                year, seq = _period_parts(label)
                execs.append({
                    "cn": cn, "period": label, "period_year": year, "period_seq": seq,
                    "hm": hm, "service_date": dt, "mekanik": mek, "status": status,
                })
                if date_raw is not None and dt is None:
                    errors.append({"file": fname, "row": r, "period": label,
                                   "error": f"tanggal tidak terbaca: {date_raw!r}"})
        return units, execs, errors
    finally:
        wb.close()


# --------------------------------------------------------------------------- import
_UNIT_UPSERT = text("""
INSERT INTO proker_unit (cn, no, unit_model, product, unit_type, hm_now, import_id)
VALUES (:cn, :no, :unit_model, :product, :unit_type, :hm_now, :import_id)
ON CONFLICT (cn) DO UPDATE SET
  no         = COALESCE(EXCLUDED.no, proker_unit.no),
  unit_model = COALESCE(EXCLUDED.unit_model, proker_unit.unit_model),
  product    = COALESCE(EXCLUDED.product, proker_unit.product),
  unit_type  = COALESCE(EXCLUDED.unit_type, proker_unit.unit_type),
  hm_now     = COALESCE(EXCLUDED.hm_now, proker_unit.hm_now),
  import_id  = EXCLUDED.import_id,
  updated_at = now()
""")

_EXEC_UPSERT = text("""
INSERT INTO proker_execution
  (cn, period, period_year, period_seq, hm, service_date, mekanik, status, import_id)
VALUES
  (:cn, :period, :period_year, :period_seq, :hm, :service_date, :mekanik, :status, :import_id)
ON CONFLICT (cn, period) DO UPDATE SET
  period_year  = EXCLUDED.period_year,
  period_seq   = EXCLUDED.period_seq,
  hm           = EXCLUDED.hm,
  service_date = EXCLUDED.service_date,
  mekanik      = EXCLUDED.mekanik,
  status       = EXCLUDED.status,
  import_id    = EXCLUDED.import_id,
  updated_at   = now()
""")


async def _existing_exec(db: AsyncSession, pairs: list[tuple[str, str]]) -> set[tuple[str, str]]:
    """Pasangan (cn, period) yang sudah ada — untuk membedakan inserted vs updated."""
    if not pairs:
        return set()
    cns = sorted({cn for cn, _ in pairs})
    rows = (await db.execute(text(
        "SELECT cn, period FROM proker_execution WHERE cn = ANY(:cns)"), {"cns": cns}
    )).mappings().all()
    return {(r["cn"], r["period"]) for r in rows} & set(pairs)


async def _run(import_id: str, files: list[tuple[str, bytes]]):
    async with SessionLocal() as db:
        try:
            merged_units: dict[str, dict] = {}
            merged_execs: dict[tuple[str, str], dict] = {}
            errors: list[dict] = []
            duplicates = 0

            for fname, content in files:
                try:
                    units, execs, errs = _parse(content, fname)
                except Exception as e:  # noqa: BLE001
                    errors.append({"file": fname, "error": str(e)})
                    continue
                errors.extend(errs[:50])
                merged_units.update(units)
                for row in execs:
                    key = (row["cn"], row["period"])
                    if key in merged_execs:
                        duplicates += 1
                    merged_execs[key] = row
                await db.execute(text(
                    "UPDATE imports SET total_rows=:t, processed_rows=:p WHERE id=:i"),
                    {"t": len(merged_execs), "p": 0, "i": import_id})
                await db.commit()

            if not merged_units:
                raise ValueError("tidak ada baris unit yang terbaca (C/N kosong semua)")

            units = list(merged_units.values())
            for i in range(0, len(units), BATCH):
                await db.execute(_UNIT_UPSERT, [
                    {**u, "import_id": import_id} for u in units[i:i + BATCH]])

            execs = list(merged_execs.values())
            inserted = updated = 0
            for i in range(0, len(execs), BATCH):
                batch = execs[i:i + BATCH]
                has = await _existing_exec(db, [(b["cn"], b["period"]) for b in batch])
                await db.execute(_EXEC_UPSERT, [
                    {**b, "import_id": import_id} for b in batch])
                updated += sum(1 for b in batch if (b["cn"], b["period"]) in has)
                inserted += sum(1 for b in batch if (b["cn"], b["period"]) not in has)
                await db.execute(text(
                    "UPDATE imports SET processed_rows=:p, ok_rows=:o WHERE id=:i"),
                    {"p": min(i + BATCH, len(execs)), "o": len(execs), "i": import_id})
                await db.commit()

            summary = {
                "units": len(units), "inserted": inserted, "updated": updated,
                "duplicates": duplicates, "files": [f for f, _ in files],
                "total_rows": len(execs), "errors": errors[:50],
            }
            await db.execute(text(
                "UPDATE imports SET status='COMMITTED', total_rows=:t, ok_rows=:o, "
                "fail_rows=:f, processed_rows=:p, errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"t": len(execs), "o": len(execs), "f": len(errors), "p": len(execs),
                 "e": json.dumps(summary, default=str), "i": import_id})
            await db.commit()
            log.info("proker import %s: %s unit, %s new, %s updated, %s dup",
                     import_id, len(units), inserted, updated, duplicates)
        except Exception as e:  # noqa: BLE001
            await db.rollback()
            await db.execute(text(
                "UPDATE imports SET status='FAILED', errors=CAST(:e AS jsonb) WHERE id=:i"),
                {"e": json.dumps({"error": str(e)}, default=str), "i": import_id})
            await db.commit()
            log.exception("proker import gagal: %s", e)


@router.post("/import", dependencies=[Depends(require_menu("proker", "add"))], status_code=202)
async def import_xlsx(
    bg: BackgroundTasks,
    files: list[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(get_current_user),
):
    """Upload 1..n file proker (.xlsx). Upsert pada (cn, period)."""
    if not files:
        raise HTTPException(400, "Tidak ada file")
    payload: list[tuple[str, bytes]] = []
    limit = settings.max_upload_mb * 1024 * 1024
    for f in files:
        content = await f.read()
        if len(content) > limit:
            raise HTTPException(413, f"{f.filename}: file terlalu besar")
        if not content:
            raise HTTPException(400, f"{f.filename}: file kosong")
        payload.append((f.filename or "upload.xlsx", content))
    label = ", ".join(name for name, _ in payload)
    res = await db.execute(text(
        "INSERT INTO imports(filename,sheet,status,total_rows,uploaded_by) "
        "VALUES (:fn,'PROKER','PROCESSING',0,:u) RETURNING id"),
        {"fn": label[:300], "u": user["username"]})
    import_id = str(res.scalar_one())
    await db.commit()
    bg.add_task(_run, import_id, payload)
    return {"import_id": import_id, "status": "PROCESSING", "files": len(payload)}


@router.get("/import/{import_id}", dependencies=[Depends(require_menu("proker", "view"))])
async def import_status(import_id: str, db: AsyncSession = Depends(get_db)):
    row = (await db.execute(text(
        "SELECT id, filename, status, total_rows, ok_rows, fail_rows, processed_rows, "
        "uploaded_by, created_at, errors FROM imports WHERE id=:i AND sheet='PROKER'"),
        {"i": import_id})).mappings().first()
    if not row:
        raise HTTPException(404, "import_id tidak ditemukan")
    return dict(row)


# --------------------------------------------------------------------------- baca
_SORT = {
    "cn": "u.cn", "no": "u.no", "unit_model": "u.unit_model", "product": "u.product",
    "unit_type": "u.unit_type", "hm_now": "u.hm_now", "exec_count": "exec_count",
    "last_date": "last_date",
}

_AGG = """
  SELECT cn,
         COUNT(*) AS exec_count,
         COUNT(*) FILTER (WHERE status = 'DONE') AS done_count,
         MAX(service_date) AS last_date,
         (ARRAY_AGG(hm ORDER BY service_date DESC NULLS LAST,
                    period_year DESC NULLS LAST, period_seq DESC NULLS LAST)
          FILTER (WHERE hm IS NOT NULL))[1] AS last_hm
  FROM proker_execution
  GROUP BY cn
"""


@router.get("/filters", dependencies=[Depends(require_menu("proker", "view"))])
async def get_filters(db: AsyncSession = Depends(get_db)):
    async def distinct(col: str) -> list[str]:
        return list((await db.execute(text(
            f"SELECT DISTINCT {col} FROM proker_unit "
            f"WHERE {col} IS NOT NULL AND {col} <> '' ORDER BY 1"))).scalars().all())

    span = (await db.execute(text(
        "SELECT MIN(service_date) AS min_date, MAX(service_date) AS max_date, "
        "MIN(period_year) AS min_year, MAX(period_year) AS max_year "
        "FROM proker_execution"))).mappings().first()
    years = list((await db.execute(text(
        "SELECT DISTINCT period_year FROM proker_execution "
        "WHERE period_year IS NOT NULL ORDER BY 1 DESC"))).scalars().all())
    return {
        "product": await distinct("product"),
        "model": await distinct("unit_model"),
        "type": await distinct("unit_type"),
        "years": years,
        "min_date": str(span["min_date"]) if span and span["min_date"] else None,
        "max_date": str(span["max_date"]) if span and span["max_date"] else None,
        "min_year": span["min_year"] if span else None,
        "max_year": span["max_year"] if span else None,
    }


@router.get("/stats", dependencies=[Depends(require_menu("proker", "view"))])
async def get_stats(db: AsyncSession = Depends(get_db)):
    unit = (await db.execute(text("SELECT COUNT(*) FROM proker_unit"))).scalar() or 0
    row = (await db.execute(text("""
        SELECT COUNT(*) AS total,
               COUNT(*) FILTER (WHERE status = 'DONE') AS done,
               COUNT(*) FILTER (WHERE status = 'PLAN') AS plan,
               COUNT(*) FILTER (WHERE status = 'LATE') AS late,
               COUNT(*) FILTER (WHERE status = 'MISSED') AS missed,
               COUNT(DISTINCT cn) FILTER (WHERE status = 'DONE') AS unit_done,
               MAX(service_date) AS last_date
        FROM proker_execution
    """))).mappings().first()
    year = datetime.date.today().year
    this_year = (await db.execute(text(
        "SELECT COUNT(*) FROM proker_execution "
        "WHERE status = 'DONE' AND EXTRACT(YEAR FROM service_date) = :y"),
        {"y": year})).scalar() or 0
    per_year = (await db.execute(text("""
        SELECT period_year AS year, COUNT(*) AS total,
               COUNT(*) FILTER (WHERE status = 'DONE') AS done
        FROM proker_execution WHERE period_year IS NOT NULL
        GROUP BY 1 ORDER BY 1
    """))).mappings().all()
    return {
        "units": unit, "total": row["total"], "done": row["done"], "plan": row["plan"],
        "late": row["late"], "missed": row["missed"], "unit_done": row["unit_done"],
        "last_date": str(row["last_date"]) if row["last_date"] else None,
        "this_year": this_year, "year": year,
        "per_year": [dict(r) for r in per_year],
    }


@router.get("/units", dependencies=[Depends(require_menu("proker", "view"))])
async def list_units(
    search: str | None = None,
    product: str | None = None,
    model: str | None = None,
    sort: str = Query("cn"),
    order: str = Query("asc", pattern="^(asc|desc)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    conds, p = [], {"lim": page_size, "off": (page - 1) * page_size}
    if search:
        conds.append("(u.cn ILIKE :s OR u.unit_model ILIKE :s OR u.product ILIKE :s "
                     "OR u.unit_type ILIKE :s)")
        p["s"] = f"%{search.strip()}%"
    if product:
        conds.append("u.product = :prod")
        p["prod"] = product
    if model:
        conds.append("u.unit_model = :model")
        p["model"] = model
    where = f"WHERE {' AND '.join(conds)}" if conds else ""
    col = _SORT.get(sort, "u.cn")
    total = (await db.execute(text(f"SELECT COUNT(*) FROM proker_unit u {where}"), p)).scalar() or 0
    rows = (await db.execute(text(f"""
        WITH e AS ({_AGG})
        SELECT u.cn, u.no, u.unit_model, u.product, u.unit_type, u.hm_now,
               COALESCE(e.exec_count, 0) AS exec_count,
               COALESCE(e.done_count, 0) AS done_count,
               e.last_date, e.last_hm
        FROM proker_unit u LEFT JOIN e ON e.cn = u.cn
        {where}
        ORDER BY {col} {order.upper()} NULLS LAST, u.cn
        LIMIT :lim OFFSET :off
    """), p)).mappings().all()
    return {"total": total, "page": page, "page_size": page_size,
            "data": [dict(r) for r in rows]}


@router.get("/units/{cn}", dependencies=[Depends(require_menu("proker", "view"))])
async def unit_detail(cn: str, db: AsyncSession = Depends(get_db)):
    key = (cn or "").strip().upper().replace(" ", "")
    unit = (await db.execute(text("SELECT * FROM proker_unit WHERE cn = :c"),
                             {"c": key})).mappings().first()
    if not unit:
        raise HTTPException(404, f"{key} tidak ditemukan")
    rows = (await db.execute(text("""
        SELECT * FROM proker_execution WHERE cn = :c
        ORDER BY service_date DESC NULLS LAST, period_year DESC NULLS LAST,
                 period_seq DESC NULLS LAST, period
    """), {"c": key})).mappings().all()
    execs = [dict(r) for r in rows]
    done = [e for e in execs if e["status"] == "DONE"]
    # HM terakhir = HM pada eksekusi selesai terbaru (server sudah urut tanggal desc).
    last = next((e for e in done if e["hm"] is not None), None)
    last_date = next((e["service_date"] for e in done if e["service_date"]), None)
    last_hm = float(last["hm"]) if last else None
    hm_now = float(unit["hm_now"]) if unit["hm_now"] is not None else None
    due = last_hm + INTERVAL_HM if last_hm is not None else None
    return {
        "unit": dict(unit),
        "executions": execs,
        "stats": {
            "total": len(execs),
            "done": len(done),
            "plan": sum(1 for e in execs if e["status"] == "PLAN"),
            "missed": sum(1 for e in execs if e["status"] == "MISSED"),
            "last_date": last_date,
            "last_hm": last_hm,
            "interval": INTERVAL_HM,
            "due_hm": due,
            "sisa_hm": (due - hm_now) if due is not None and hm_now is not None else None,
        },
    }


# --------------------------------------------------------------------------- tulis
@router.post("/units", dependencies=[Depends(require_menu("proker", "add"))], status_code=201)
async def create_unit(body: dict, db: AsyncSession = Depends(get_db)):
    cn = (_txt(body.get("cn")) or "").upper().replace(" ", "")
    if not cn:
        raise HTTPException(400, "cn wajib diisi")
    if (await db.execute(text("SELECT 1 FROM proker_unit WHERE cn=:c"), {"c": cn})).scalar():
        raise HTTPException(409, f"{cn} sudah ada")
    row = (await db.execute(text("""
        INSERT INTO proker_unit (cn, no, unit_model, product, unit_type, hm_now)
        VALUES (:cn, :no, :unit_model, :product, :unit_type, :hm_now) RETURNING *
    """), {
        "cn": cn, "no": _int(body.get("no")),
        "unit_model": _txt(body.get("unit_model")), "product": _txt(body.get("product")),
        "unit_type": _txt(body.get("unit_type")), "hm_now": _num(body.get("hm_now")),
    })).mappings().first()
    await db.commit()
    return dict(row)


@router.patch("/units/{cn}", dependencies=[Depends(require_menu("proker", "edit"))])
async def update_unit(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    key = cn.strip().upper()
    sets, params = [], {"c": key}
    for f in UNIT_FIELDS:
        if f not in body:
            continue
        params[f] = _num(body[f]) if f == "hm_now" else (
            _int(body[f]) if f == "no" else _txt(body[f]))
        sets.append(f"{f}=:{f}")
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    res = await db.execute(text(
        f"UPDATE proker_unit SET {', '.join(sets)} WHERE cn=:c"), params)
    if res.rowcount == 0:
        raise HTTPException(404, f"{key} tidak ditemukan")
    await db.commit()
    row = (await db.execute(text("SELECT * FROM proker_unit WHERE cn=:c"),
                            {"c": key})).mappings().first()
    return dict(row)


@router.post("/units/{cn}/executions",
             dependencies=[Depends(require_menu("proker", "add"))], status_code=201)
async def create_execution(cn: str, body: dict, db: AsyncSession = Depends(get_db)):
    key = cn.strip().upper()
    if not (await db.execute(text("SELECT 1 FROM proker_unit WHERE cn=:c"), {"c": key})).scalar():
        raise HTTPException(404, f"{key} belum ada di master unit")
    period = _txt(body.get("period"))
    if not period:
        raise HTTPException(400, "periode wajib diisi")
    hm = _num(body.get("hm"))
    year, seq = _period_parts(period)
    dup = (await db.execute(text(
        "SELECT 1 FROM proker_execution WHERE cn=:c AND period=:p"),
        {"c": key, "p": period})).scalar()
    if dup:
        raise HTTPException(409, f"periode {period} untuk {key} sudah ada")
    row = (await db.execute(text("""
        INSERT INTO proker_execution
          (cn, period, period_year, period_seq, hm, service_date, mekanik, status)
        VALUES (:cn, :period, :y, :seq, :hm, :d, :mek, :status) RETURNING *
    """), {
        "cn": key, "period": period, "y": year, "seq": seq, "hm": hm,
        "d": _date(body.get("service_date")), "mek": _txt(body.get("mekanik")),
        "status": _norm_status(body.get("status"), hm),
    })).mappings().first()
    await db.commit()
    return dict(row)


@router.patch("/executions/{exec_id}", dependencies=[Depends(require_menu("proker", "edit"))])
async def update_execution(exec_id: int, body: dict, db: AsyncSession = Depends(get_db)):
    cur = (await db.execute(text("SELECT * FROM proker_execution WHERE id=:i"),
                            {"i": exec_id})).mappings().first()
    if not cur:
        raise HTTPException(404, "baris proker tidak ditemukan")

    sets, params = [], {"i": exec_id}
    period = _txt(body.get("period"))
    if period and period != cur["period"]:
        dup = (await db.execute(text(
            "SELECT 1 FROM proker_execution WHERE cn=:c AND period=:p AND id<>:i"),
            {"c": cur["cn"], "p": period, "i": exec_id})).scalar()
        if dup:
            raise HTTPException(409, f"periode {period} untuk {cur['cn']} sudah ada")
        year, seq = _period_parts(period)
        sets += ["period=:period", "period_year=:y", "period_seq=:seq"]
        params.update({"period": period, "y": year, "seq": seq})
    if "hm" in body:
        sets.append("hm=:hm")
        params["hm"] = _num(body["hm"])
    if "service_date" in body:
        sets.append("service_date=:d")
        params["d"] = _date(body["service_date"])
    if "mekanik" in body:
        sets.append("mekanik=:mek")
        params["mek"] = _txt(body["mekanik"])
    if "status" in body:
        hm_val = params.get("hm") if "hm" in params else _num(cur["hm"])
        sets.append("status=:status")
        params["status"] = _norm_status(body["status"], hm_val)
    if not sets:
        raise HTTPException(400, "tidak ada field yang diubah")
    row = (await db.execute(text(
        f"UPDATE proker_execution SET {', '.join(sets)} WHERE id=:i RETURNING *"),
        params)).mappings().first()
    await db.commit()
    return dict(row)


@router.delete("/executions/{exec_id}",
               dependencies=[Depends(require_menu("proker", "delete"))])
async def delete_execution(exec_id: int, db: AsyncSession = Depends(get_db)):
    res = await db.execute(text("DELETE FROM proker_execution WHERE id=:i"), {"i": exec_id})
    if res.rowcount == 0:
        raise HTTPException(404, "baris proker tidak ditemukan")
    await db.commit()
    return {"deleted": res.rowcount}
